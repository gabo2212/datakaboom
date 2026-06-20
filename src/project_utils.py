from __future__ import annotations

import csv
import datetime as dt
import math
import os
import re
import unicodedata
from collections import Counter
from pathlib import Path
from typing import Iterable, Iterator

import pandas as pd
from openpyxl import load_workbook
from rapidfuzz import fuzz


ROOT_DIR = Path(__file__).resolve().parent.parent
DATA_DIR = ROOT_DIR / "data"
RAW_DIR = DATA_DIR / "raw"
PROCESSED_DIR = DATA_DIR / "processed"
REPORTS_DIR = ROOT_DIR / "reports"
SRC_DIR = ROOT_DIR / "src"

DEFAULT_SOURCE_SHEET = "Extraction brute"
DEFAULT_RAW_CSV = RAW_DIR / "fichier_original.csv"
DEFAULT_PROCESSED_CSV = PROCESSED_DIR / "fichier_nettoye.csv"
DEFAULT_EDA_REPORT = REPORTS_DIR / "eda_report.md"
DEFAULT_MISSING_VALUES_REPORT = REPORTS_DIR / "missing_values_report.csv"
DEFAULT_DUPLICATES_REPORT = REPORTS_DIR / "duplicates_report.csv"
DEFAULT_VARIANTS_REPORT = REPORTS_DIR / "organization_variants_report.csv"
DEFAULT_MAPPING_REPORT = REPORTS_DIR / "entity_cleaning_mapping.csv"
DEFAULT_CLEANING_EXECUTION_REPORT = REPORTS_DIR / "cleaning_execution_report.md"
DEFAULT_FINAL_SUMMARY = REPORTS_DIR / "final_summary.md"


def ensure_output_directories() -> None:
    """Create the project folders expected by the assignment."""

    RAW_DIR.mkdir(parents=True, exist_ok=True)
    PROCESSED_DIR.mkdir(parents=True, exist_ok=True)
    REPORTS_DIR.mkdir(parents=True, exist_ok=True)
    SRC_DIR.mkdir(parents=True, exist_ok=True)


def candidate_workbook_paths() -> list[Path]:
    """Return workbook candidates in priority order."""

    candidates = [
        RAW_DIR / "fichier_original.xlsx",
        DATA_DIR
        / "2026-05-13_donnees-ouvertes_divulgation-octrois-subventions-et-contributions.xlsx",
    ]
    # Fallback for alternate naming if the workbook was renamed.
    candidates.extend(sorted(DATA_DIR.glob("*.xlsx")))
    unique: list[Path] = []
    seen: set[Path] = set()
    for candidate in candidates:
        if candidate in seen:
            continue
        seen.add(candidate)
        unique.append(candidate)
    return unique


def discover_workbook_path() -> Path | None:
    """Find the Excel workbook containing the raw data."""

    for candidate in candidate_workbook_paths():
        if candidate.exists():
            return candidate
    return None


def ensure_raw_csv(
    *,
    source_sheet: str = DEFAULT_SOURCE_SHEET,
    force_refresh: bool = False,
) -> Path:
    """
    Ensure the CSV snapshot used by the pipeline exists.

    The assignment talks about CSV files, but the provided source is an Excel
    workbook. We materialize the raw sheet into `data/raw/fichier_original.csv`
    so the rest of the pipeline can use chunked CSV processing.
    """

    ensure_output_directories()

    if DEFAULT_RAW_CSV.exists() and not force_refresh:
        return DEFAULT_RAW_CSV

    workbook_path = discover_workbook_path()
    if workbook_path is None:
        raise FileNotFoundError(
            "Unable to find the source workbook. Expected one of: "
            + ", ".join(str(path) for path in candidate_workbook_paths())
        )

    extract_workbook_sheet_to_csv(
        workbook_path=workbook_path,
        csv_path=DEFAULT_RAW_CSV,
        sheet_name=source_sheet,
    )
    return DEFAULT_RAW_CSV


def extract_workbook_sheet_to_csv(
    workbook_path: Path,
    csv_path: Path,
    sheet_name: str,
) -> None:
    """Stream one worksheet to CSV without loading the workbook in memory."""

    if not workbook_path.exists():
        raise FileNotFoundError(workbook_path)

    workbook = load_workbook(workbook_path, read_only=True, data_only=True)
    try:
        if sheet_name not in workbook.sheetnames:
            available = ", ".join(workbook.sheetnames)
            raise ValueError(
                f"Sheet '{sheet_name}' was not found in {workbook_path.name}. "
                f"Available sheets: {available}"
            )

        worksheet = workbook[sheet_name]
        with csv_path.open("w", encoding="utf-8", newline="") as handle:
            writer = csv.writer(handle, lineterminator="\n")
            for row in worksheet.iter_rows(values_only=True):
                writer.writerow([excel_value_to_text(value) for value in row])
    finally:
        workbook.close()


def excel_value_to_text(value: object) -> str:
    """Convert workbook cell values to stable CSV text."""

    if value is None:
        return ""

    if isinstance(value, pd.Timestamp):
        return value.isoformat(sep=" ")
    if isinstance(value, dt.datetime):
        if value.time() == dt.time(0, 0):
            return value.date().isoformat()
        return value.isoformat(sep=" ")
    if isinstance(value, dt.date):
        return value.isoformat()

    return str(value)


def read_csv_sample(
    csv_path: Path,
    *,
    nrows: int = 1000,
    usecols: list[str] | None = None,
) -> pd.DataFrame:
    """Read a small sample from the CSV snapshot."""

    return pd.read_csv(
        csv_path,
        nrows=nrows,
        usecols=usecols,
        dtype=str,
        keep_default_na=True,
        low_memory=False,
    )


def iter_csv_chunks(
    csv_path: Path,
    *,
    chunksize: int = 50_000,
    usecols: list[str] | None = None,
) -> Iterator[pd.DataFrame]:
    """Yield CSV chunks as string-based DataFrames."""

    return pd.read_csv(
        csv_path,
        chunksize=chunksize,
        usecols=usecols,
        dtype=str,
        keep_default_na=True,
        low_memory=False,
    )


def count_csv_rows(csv_path: Path) -> int:
    """Count parsed data rows in a CSV file without loading it fully."""

    total = 0
    for chunk in pd.read_csv(
        csv_path,
        usecols=[0],
        chunksize=50_000,
        dtype=str,
        keep_default_na=True,
        low_memory=False,
    ):
        total += len(chunk)
    return total


def file_size_bytes(path: Path) -> int:
    return path.stat().st_size


def human_readable_size(num_bytes: int) -> str:
    units = ["B", "KB", "MB", "GB", "TB"]
    size = float(num_bytes)
    for unit in units:
        if size < 1024 or unit == units[-1]:
            if unit == "B":
                return f"{int(size)} {unit}"
            return f"{size:.2f} {unit}"
        size /= 1024
    return f"{size:.2f} TB"


def is_missing_value(value: object) -> bool:
    if value is None:
        return True
    try:
        if pd.isna(value):
            return True
    except Exception:
        pass
    return False


def normalize_text(value: object) -> str:
    """Normalize text the way the assignment expects."""

    if is_missing_value(value):
        return ""

    text = str(value).strip().upper()
    text = unicodedata.normalize("NFKD", text)
    text = "".join(char for char in text if not unicodedata.combining(char))
    text = re.sub(r"[^A-Z0-9 ]", " ", text)
    text = re.sub(r"\s+", " ", text)
    return text.strip()


def compact_text(value: object) -> str:
    return normalize_text(value).replace(" ", "")


def collapse_repeated_segments(value: object) -> str:
    """
    Collapse values that repeat the same token separated by / or |.

    This catches cases like:
    - "Université du Québec à Montréal|Université du Québec à Montréal"
    - "X / X"
    """

    if is_missing_value(value):
        return ""

    text = str(value).strip()
    if not text:
        return ""

    parts = [part.strip() for part in re.split(r"[|/]", text) if part.strip()]
    if len(parts) <= 1:
        return text

    normalized_parts = [normalize_text(part) for part in parts]
    if len(set(normalized_parts)) == 1:
        return parts[0]

    return text


UQAM_ACCEPT_NORMALIZED = {
    "UQAM",
    "U Q A M",
    "UNIVERSITE DU QUEBEC A MONTREAL",
}

UQAM_REVIEW_NORMALIZED = {
    "AQAM",
    "QAM",
    "COOP UQAM",
}


def is_uqam_accept_alias(value: object) -> bool:
    normalized = normalize_text(collapse_repeated_segments(value))
    return normalized in UQAM_ACCEPT_NORMALIZED


def is_uqam_review_candidate(value: object) -> bool:
    normalized = normalize_text(collapse_repeated_segments(value))
    compact = normalized.replace(" ", "")
    if not normalized:
        return False

    if normalized in UQAM_REVIEW_NORMALIZED:
        return True

    # Short acronym-like strings that contain QAM/UQAM but are not the exact
    # approved aliases are suspicious enough to keep in review.
    if len(compact) <= 12 and ("QAM" in compact or "UQAM" in compact):
        return not is_uqam_accept_alias(value)

    return False


def fuzzy_similarity_score(source: object, target: object) -> int:
    """
    Compute a conservative similarity score.

    We average the standard ratio and the partial ratio to avoid over-boosting
    substring matches such as "COOP UQAM".
    """

    source_compact = compact_text(source)
    target_compact = compact_text(target)
    if not source_compact or not target_compact:
        return 0

    ratio_score = fuzz.ratio(source_compact, target_compact)
    partial_score = fuzz.partial_ratio(source_compact, target_compact)
    return int(round((ratio_score + partial_score) / 2))


def infer_column_roles(sample: pd.DataFrame) -> dict[str, list[str]]:
    """Infer numeric, text, and date-like columns from a sample."""

    numeric_columns: list[str] = []
    text_columns: list[str] = []
    date_columns: list[str] = []

    for column in sample.columns:
        series = sample[column].dropna().astype(str).str.strip()
        series = series[series != ""]
        if series.empty:
            text_columns.append(column)
            continue

        numeric_rate = _numeric_parse_rate(series)
        column_name = column.lower()
        date_rate = _date_parse_rate(series) if _looks_date_candidate(column_name) else 0.0

        if _looks_date_like(column_name, date_rate, series):
            date_columns.append(column)
        elif _looks_numeric_like(column_name, numeric_rate, series):
            numeric_columns.append(column)
        else:
            text_columns.append(column)

    return {
        "numeric": numeric_columns,
        "text": text_columns,
        "date": date_columns,
    }


def _date_parse_rate(series: pd.Series) -> float:
    text = series.astype(str).str.strip()
    pattern = (
        r"^\d{4}-\d{2}-\d{2}$"
        r"|^\d{4}-\d{2}-\d{2}\s+\d{2}:\d{2}:\d{2}$"
        r"|^\d{4}/\d{2}/\d{2}$"
        r"|^\d{1,2}/\d{1,2}/\d{4}$"
    )
    matched = text.str.match(pattern, na=False)
    return float(matched.mean())


def _looks_date_candidate(column_name: str) -> bool:
    return any(token in column_name for token in ("date", "start", "end"))


def _numeric_parse_rate(series: pd.Series) -> float:
    cleaned = series.str.replace(",", "", regex=False).str.replace("$", "", regex=False)
    cleaned = cleaned.str.replace(" ", "", regex=False)
    parsed = pd.to_numeric(cleaned, errors="coerce")
    return float(parsed.notna().mean())


def _looks_date_like(column_name: str, date_rate: float, series: pd.Series) -> bool:
    if "date" in column_name:
        return date_rate >= 0.5
    if date_rate >= 0.9:
        # Avoid classifying short numeric identifiers as dates.
        sample_lengths = series.astype(str).str.len().median()
        return sample_lengths >= 8
    return False


def _looks_numeric_like(column_name: str, numeric_rate: float, series: pd.Series) -> bool:
    if numeric_rate < 0.9:
        return False
    if "name" in column_name or "title" in column_name or "description" in column_name:
        return False
    median_length = series.astype(str).str.len().median()
    return median_length <= 20


def dataframe_to_markdown(df: pd.DataFrame, *, index: bool = False) -> str:
    """Render a dataframe as GitHub-flavored markdown."""

    try:
        return df.to_markdown(index=index)
    except Exception:
        # Fallback if tabulate or rich markdown rendering is unavailable.
        return df.to_string(index=index)


def write_text(path: Path, text: str) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(text, encoding="utf-8")


def write_dataframe(path: Path, df: pd.DataFrame) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    df.to_csv(path, index=False, encoding="utf-8")


def read_dataframe(path: Path, **kwargs) -> pd.DataFrame:
    return pd.read_csv(path, dtype=str, keep_default_na=True, low_memory=False, **kwargs)


def choose_canonical_label(values: Iterable[str], counts: Counter[str]) -> str:
    """Pick the most useful canonical label from a compact group."""

    candidates = [value for value in values if not is_missing_value(value)]
    if not candidates:
        return ""

    def sort_key(value: str) -> tuple[int, int, str]:
        normalized = normalize_text(value)
        return (-counts.get(value, 0), len(normalized), normalized)

    return sorted(candidates, key=sort_key)[0]
